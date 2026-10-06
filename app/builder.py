"""Upload and visual-builder endpoints. All storage names are server generated."""
import csv
import io
import json
import re
import shutil
import warnings
import zipfile
from datetime import date, datetime
from pathlib import Path
from uuid import UUID

from fastapi import HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from PIL import Image
from pydantic import BaseModel, Field, model_validator

from app.database import Item, Job, identifier, now

FONTS = {"Helvetica", "Helvetica-Bold", "Times-Roman", "Times-Bold", "Courier"}


class TextField(BaseModel):
    column: str = Field(min_length=1, max_length=100)
    x_mm: float = Field(ge=0, le=420)
    y_mm: float = Field(ge=0, le=420)
    max_width_mm: float = Field(gt=0, le=420)
    font: str = "Helvetica"
    font_size_pt: int = Field(default=24, ge=6, le=96)
    min_font_size_pt: int = Field(default=10, ge=6, le=96)
    color: str = "#203040"
    align: str = "center"

    @model_validator(mode="after")
    def valid(self):
        if self.font not in FONTS or self.align not in {"center", "left", "right"}:
            raise ValueError("Unsupported font or alignment.")
        if not re.fullmatch(r"#[0-9a-fA-F]{6}", self.color):
            raise ValueError("Color must be a six-digit hex color.")
        if self.min_font_size_pt > self.font_size_pt:
            raise ValueError("Minimum font size exceeds font size.")
        return self


class DesignRequest(BaseModel):
    image_id: UUID
    sheet_id: UUID
    name_column: str
    title: str = Field(default="Custom certificates", min_length=1, max_length=150)
    width_mm: float = Field(default=297, ge=50, le=420)
    height_mm: float = Field(default=210, ge=50, le=420)
    fields: list[TextField] = Field(min_length=1, max_length=30)


def install_builder(app, settings):
    static = Path(__file__).parent / "static"
    app.mount("/static", StaticFiles(directory=static), name="static")

    @app.get("/", include_in_schema=False)
    def home():
        return FileResponse(static / "index.html")

    def location(value, suffix):
        return settings.storage_dir / "uploads" / f"{value}{suffix}"

    async def read_upload(file):
        data = await file.read(10 * 1024 * 1024 + 1)
        await file.close()
        if len(data) > 10 * 1024 * 1024:
            raise HTTPException(413, "Upload must be at most 10 MB.")
        return data

    @app.post("/builder/images")
    async def upload_image(file: UploadFile):
        data = await read_upload(file)
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", Image.DecompressionBombWarning)
                with Image.open(io.BytesIO(data)) as image:
                    if image.format not in {"PNG", "JPEG"}:
                        raise ValueError("Unsupported image format")
                    width, height = image.size
                    if width * height > 20_000_000:
                        raise ValueError("Image exceeds 20 megapixels")
                    image.load()
                    image_id = identifier()
                    target = location(image_id, ".png")
                    target.parent.mkdir(parents=True, exist_ok=True)
                    image.convert("RGB").save(target)
        except Exception as exc:
            raise HTTPException(422, "Upload a valid PNG or JPEG, at most 20 megapixels.") from exc
        return {"image_id": image_id, "width": width, "height": height,
                "url": f"/builder/images/{image_id}"}

    @app.get("/builder/images/{image_id}")
    def image(image_id: UUID):
        target = location(image_id, ".png")
        if not target.exists():
            raise HTTPException(404, "Image not found.")
        return FileResponse(target, media_type="image/png")

    @app.post("/builder/sheets")
    async def upload_sheet(file: UploadFile):
        filename = file.filename or ""
        data = await read_upload(file)
        try:
            if filename.lower().endswith(".csv"):
                iterator = iter(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
                workbook = None
            elif filename.lower().endswith(".xlsx"):
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                        raise ValueError("Expanded workbook must be at most 50 MB")
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
                iterator = workbook.active.iter_rows(values_only=True)
            else:
                raise ValueError("Use .xlsx or .csv")
            try:
                header = next(iterator)
                if len(header) > 50:
                    raise ValueError("At most 50 columns are allowed")
                columns = [str(value).strip() if value is not None else "" for value in header]
                if not columns or any(not value or len(value) > 100 for value in columns) or len(set(columns)) != len(columns):
                    raise ValueError("First row must contain unique, non-empty column names (up to 100 characters)")
                rows = []
                for values in iterator:
                    if all(value is None or value == "" for value in values):
                        continue
                    if len(rows) >= settings.max_recipients:
                        raise ValueError(f"At most {settings.max_recipients} data rows are allowed")
                    if len(values) > len(columns) and any(v not in (None, "") for v in values[len(columns):]):
                        raise ValueError("Data contains cells without a column heading")
                    row = {}
                    for index, column in enumerate(columns):
                        value = values[index] if index < len(values) else None
                        text = value.isoformat() if isinstance(value, (date, datetime)) else str(value) if value is not None else ""
                        if len(text) > 500:
                            raise ValueError("Cell values must be at most 500 characters")
                        row[column] = text
                    rows.append(row)
                if not rows:
                    raise ValueError("Sheet has no data rows")
            finally:
                if workbook:
                    workbook.close()
        except Exception as exc:
            raise HTTPException(422, f"Could not read spreadsheet: {exc}") from exc
        sheet_id = identifier()
        target = location(sheet_id, ".json")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(json.dumps({"columns": columns, "rows": rows}), encoding="utf-8")
        return {"sheet_id": sheet_id, "columns": columns, "rows": rows[:10], "total": len(rows)}

    @app.post("/builder/jobs", status_code=202)
    def generate(payload: DesignRequest):
        image_path = location(payload.image_id, ".png")
        sheet_path = location(payload.sheet_id, ".json")
        if not image_path.exists() or not sheet_path.exists():
            raise HTTPException(404, "Upload the image and spreadsheet first.")
        sheet = json.loads(sheet_path.read_text(encoding="utf-8"))
        if payload.name_column not in sheet["columns"] or any(field.column not in sheet["columns"] for field in payload.fields):
            raise HTTPException(422, "A mapped column does not exist in the spreadsheet.")
        for field in payload.fields:
            if field.x_mm > payload.width_mm or field.y_mm > payload.height_mm:
                raise HTTPException(422, "Field position lies outside the page.")
            left = field.x_mm - (field.max_width_mm / 2 if field.align == "center" else field.max_width_mm if field.align == "right" else 0)
            if left < -.01 or left + field.max_width_mm > payload.width_mm + .01:
                raise HTTPException(422, "Field text width lies outside the page.")
        job_id = identifier()
        directory = settings.storage_dir / "designs" / job_id
        directory.mkdir(parents=True)
        try:
            shutil.copyfile(image_path, directory / "background.png")
            layout = {"version": "custom-1", "page": {"width_mm": payload.width_mm, "height_mm": payload.height_mm},
                      "fields": {str(i): field.model_dump(exclude={"column"}) for i, field in enumerate(payload.fields)}}
            (directory / "layout.json").write_text(json.dumps(layout), encoding="utf-8")
            values_by_id = {}
            with app.state.sessions.begin() as session:
                job = Job(id=job_id, course_name=payload.title, issue_date=date.today().isoformat(),
                          template_version="custom-1", total_count=len(sheet["rows"]))
                session.add(job)
                session.flush()
                accepted = 0
                for index, row in enumerate(sheet["rows"]):
                    name = row[payload.name_column].strip()
                    values = {str(i): row[field.column] for i, field in enumerate(payload.fields)}
                    valid = bool(name) and len(name) <= 150 and all(all(32 <= ord(c) <= 126 for c in text) for text in [name, *values.values()])
                    item = Item(job_id=job_id, input_index=index, recipient_name=name, status="pending" if valid else "failed")
                    if valid:
                        item.certificate_id = identifier()
                        values_by_id[item.certificate_id] = values
                        accepted += 1
                    else:
                        item.error_stage = "validation"
                        item.error_code = "INVALID_RECIPIENT"
                        item.error_message = "Name must be non-empty (up to 150 characters); mapped text must use printable ASCII."
                        item.completed_at = now()
                    session.add(item)
                (directory / "values.json").write_text(json.dumps(values_by_id), encoding="utf-8")
                if not accepted:
                    job.status = "failed"
                    job.completed_at = now()
            return {"job_id": job_id, "accepted": accepted, "rejected": job.total_count - accepted, "status": job.status}
        except BaseException:
            shutil.rmtree(directory)
            raise
